import tkinter as tk
import tkinter.messagebox as tkMessageBox
import logging
import re
import openpyxl
import excelxml

import mywidgets.userinterface as userinterface
from mywidgets.Tools.mywinzip.file_menu import FileMenu
from mywidgets.equations import equations_manager
from worksheetui import SheetState, SheetUI, test_content_gen, SheetLook

logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

def alpha_code(id, nbase=26):
    answ = []
    while id:
        res = id % nbase
        id = id // nbase
        if res == 0:
            res = nbase
            id -=1
        answ.append(chr(ord('A') + res - 1))
    return ''.join(answ[::-1])



class WbViewer(tk.Tk):

    wb_manager = excelxml
    
    def __init__(self):
        super().__init__()
        # self.sheetui: SheetUI
        self.event_add('<<MENUCLICK>>', 'None')
        self.event_add('<<VAR_CHANGE>>', 'None')
        self.bind_all('<<MENUCLICK>>', self.onMenuClick)
        self.bind_all('<<VAR_CHANGE>>', self.onVarChange)
        self.setGUI()

        self.fmngr = fmngr = FileMenu(self)
        fmngr.default_file_name = 'newZip.zip'
        fmngr.default_path = 'C:/Users/agmontesb/Downloads/'
        fmngr.default_extension = '.xlsx'
        # fmngr.default_file_type = [('Excel Workbook', '*.xlsx'), ('Excel 97-2003 Workbook', '*.xls')]
        fmngr.default_file_type = ('Excel Workbook', '*.xlsx')

        self.ws_ctxs = {}
        self.named_range = {}
        self.active_sheet = self.cb_sheet_selector.get()
        self._r1c1_flag = True
        # self.r1c1_flag = False

        self.wb = None
        pass

    def content_gen(self, nquadrant, x, y):
        if self.wb is None:
            return test_content_gen(nquadrant, x, y)
        sheet = self.wb[self.active_sheet]
        return sheet.cell(row=y, column=x).value


    @property
    def r1c1_flag(self) -> bool:
        return self._r1c1_flag
    
    @r1c1_flag.setter
    def r1c1_flag(self, value: bool):
        if self._r1c1_flag != value:
            if not value:
                self.sheetui.format_header = lambda index, axis: alpha_code(index) if axis == 1 else f"{index}"
            else:
                self.sheetui.format_header = None
        self._r1c1_flag = value

    def on_active_cell_changed(wb, event):
        wdg: SheetUI = event.widget
        if wb.wb:
            cell = wb.wb[wb.active_sheet].cell(*wdg.active_cell[::-1])
            try:
                item = cell.formula
            except AttributeError:
                item = cell.value
        else:
            item = test_content_gen(0, *wdg.active_cell[::-1])
        wb.lbl_cell_content['text'] = str(item)

    def on_selected_cells_changed(self, event):
        wdg: SheetUI = event.widget
        selected_cells = wdg.selected_cells
        sel_x0, sel_y0, sel_x1, sel_y1 = selected_cells
        if (sel_x0, sel_y0) == (sel_x1, sel_y1):
            msg = self.sheetui.format_header(sel_x0, axis=1) + self.sheetui.format_header(sel_y0, axis=0)
        else:
            nrows = sel_y1 - sel_y0 + 1
            ncols = sel_x1 - sel_x0 + 1
            msg = f'{nrows}R x {ncols}C'
        self.cb_named_range.set(msg)

    def setGUI(self):
        file_path = '@layout/wbviewerui'
        xmlObj = userinterface.getLayout(file_path, withCss=False)
        userinterface.newPanelFactory(
            master=self,
            selpane=xmlObj,
            genPanelModule=None,
            setParentTo='master',
            registerWidget=self.register_widget,
        )
        equations_manager.set_initial_widget_states()

        try:
            cb_sheet_selector = self.cb_sheet_selector
            cb_sheet_selector.bind("<<ComboboxSelected>>", self.on_combobox_change)
        except AttributeError:
            pass
        cb_sheet_selector.current(0)

        try:
            cb_named_range = self.cb_named_range
            cb_named_range.bind("<Return>", self.on_cbrange_return)
            cb_named_range.bind("<FocusIn>", self.on_cbrange_focusin)
            cb_named_range.bind("<<ComboboxSelected>>", self.on_cbrange_change)

        except AttributeError:
            pass

        try:
            sheetui:SheetUI = self.sheetui
            sheetui.focus_set()
            if logger.isEnabledFor(logging.DEBUG):
                sheetui.cell_content = self.content_gen
        except AttributeError:
            pass

        self.bind("<<SelectedCellsChanged>>", self.on_selected_cells_changed)
        self.bind("<<ActiveCellChanged>>", self.on_active_cell_changed)
        pass

    def on_cbrange_focusin(self, event):
        """Selects all text in the combobox when clicked or focused."""
        wdg = event.widget
        self.after(10, lambda: wdg.selection_range(0, 'end'))

    def on_cbrange_change(self, event):
        widget = event.widget
        range_name = widget.get()
        active_sheet, coords = self.named_range[range_name]
        if self.active_sheet != active_sheet:
            self.cb_sheet_selector.set(active_sheet)
            self.cb_sheet_selector.event_generate("<<ComboboxSelected>>")
        self.sheetui.set_selected_cells(*coords)
        self.sheetui.focus_set()

    def on_cbrange_return(self, event):
        """Adds the typed text to the combobox values when Return is pressed."""
        widget = event.widget
        new_value = widget.get()

        # Check for range format R1C1:R3C3
        re_str = r'R(\d+)C(\d+)(?::R(\d+)C(\d+))*' if self.r1c1_flag else r'([A-Z]+)(\d+)(?::([A-Z]+)(\d+))*'
        m = re.match(re_str, new_value.upper())
        if m:
            coords = m.groups()
            if ':' not in new_value:
                coords = coords[:2]
            if self.r1c1_flag:
                # El método set_selected_cells espera col, row, col, row y m.groups() devuelve row, col, row, col
                tlcorner, brcorner = coords[:2], coords[2:]
                coords = (*tlcorner[::-1], *brcorner[::-1])

            fnc = lambda x: int(x) if x.isdigit() else sum((ord(c) - ord('A') + 1) * (26 ** i) for i, c in enumerate(reversed(x)))
            self.sheetui.set_selected_cells(*map(fnc, coords))
            self.sheetui.focus_set()
            self.cb_named_range.set(new_value.upper())
            return

        # Se quiere agregar un nuevo nombre de rango

        # Get the current list of values from the combobox
        current_values = list(widget["values"])

        if new_value not in current_values:
            current_values.append(new_value)
            widget.config(values=sorted(current_values))
            # The new value is already displayed as it was typed by the user.
            self.named_range[new_value] = (self.active_sheet, self.sheetui.selected_cells)
            widget.set(new_value)
            self.sheetui.focus_set()
        else:
            self.cb_named_range.set(new_value)
            self.cb_named_range.event_generate("<<ComboboxSelected>>")

    def on_combobox_change(self, event):
        wdg = event.widget
        active_sheet = self.active_sheet
        self.ws_ctxs[active_sheet] = self.sheetui.look

        self.active_sheet = ws_name = wdg.get()
        look = self.ws_ctxs.get(ws_name, SheetLook())
        self.sheetui.set_sheet(look)
        self.sheetui.focus_set()
        pass

    def onVarChange(self, event=None, attr_data=None):
        if event:
            attr_data = event.attr_data
        var_name, value = attr_data
        match var_name:
            case 'view_headings':
                self.sheetui.toggle_headings()
            case 'view_gridlines':
                self.sheetui.toggle_gridlines()
            case 'fml_r1c1':
                self.r1c1_flag = value
                self.sheetui.redraw_headings()
            case 'wb_loader':
                if value == 'openpyxl':
                    self.wb_manager = openpyxl
                else:
                    self.wb_manager = excelxml

    def onMenuClick(self, event):
        menu_master, indx = event.widget, event.data
        menu, menu_item = menu_master.cget("title"), menu_master.entrycget(indx, "label")
        match menu:
            case 'file' | 'open recent':
                self.file_menu(menu_master, indx)
            case 'view':
                self.view_menu(menu_master, indx)

    def view_menu(self, menu_master: tk.Menu, indx: int):
        logger.debug(f"View menu item selected: {menu_item}")
        menu_item: str = menu_master.entrycget(indx, "label")
        sheetui: SheetUI = self.sheetui
        match menu_item:
            case x if x.startswith('Freeze'):
                menu_item = menu_item.split(' ', 1)[1]
                if menu_item == 'Panes':
                    sheetui.toggle_freeze_panes()
                else:
                    bflag = menu_item == 'First Column'
                    nq = sheetui.cell_quadrant(*sheetui.coords_vportq3)
                    vx, vy = sheetui.quadrant_origin(nq, False)
                    if sheetui.flags & SheetState.FREEZE:
                        sheetui.unfreeze_panes()
                    cx, cy = sheetui.quadrant_origin(1, False)
                    dx, dy = (1, 0) if bflag else (0, 1)
                    sheetui.freeze_panes(cx + dx, cy + dy)
                    sheetui.move_viewport(vx + dx, vy + dy)
                    sheetui.show_ws_elements()
                    indx -= 1 if menu_item == 'Top Row' else 2
                menu_master.entryconfig(indx, label='UnFreeze Panes')
            case _: # 'UnFreeze Panes'
                assert menu_item == 'UnFreeze Panes'
                sheetui.unfreeze_panes()
                menu_master.entryconfig(indx, label='Freeze Panes')


    def file_menu(self, menu_master: tk.Menu, indx: int):
        menu, menu_item = menu_master.cget("title"), menu_master.entrycget(indx, "label")
        logger.debug(f"Menu '{menu}' item selected: '{menu_item}'")
        if menu == 'file':
            match menu_item:
                case 'Open':
                    with self.fmngr.openFile() as filename:
                        assert filename
                        self.loadwb(filename)
                    pass
                case _:
                    pass
        elif menu == 'open recent':
            indx, filename = menu_item.split()
            wbfilename = self.fmngr.fileHistory[int(indx) - 1]
            assert wbfilename.endswith(filename)
            self.loadwb(wbfilename, mode='a')
    
    def loadwb(self, filename):
        load_workbook = getattr(self.wb_manager, 'load_workbook')
        try:
            self.wb = wb = load_workbook(filename)
            sheet_names = wb.sheetnames
            self.cb_sheet_selector['values'] = sheet_names

            self.active_sheet = active_sheet = wb.active.title
            self.cb_sheet_selector.set(active_sheet)
            self.ws_ctxs = {}
            self.named_range = {}
            self.sheetui.set_sheet()

        except Exception as e:
            tkMessageBox.showerror(title='Loading Error', message=str(e))
        self.fmngr.fileHistory = self.fmngr.recFile(filename)
        self.fmngr.title(filename)

    def registerMenu(self, parent, selPane, menu_master, labels):
        title = menu_master.cget('title')
        match title:
            case 'open recent':
                menu_master.config(postcommand=lambda: self.fmngr.fileHist(menu_master))

    def register_widget(self, master, xmlwidget, widget):
        attribs = xmlwidget.attrib
        name = attribs.get('name')
        if name in ('cb_sheet_selector', 'cb_named_range', 'lbl_cell_content', 'sheetui'):
            setattr(self, name, widget)
        pass


def main():
    wbv = WbViewer()
    wbv.mainloop()


if __name__ == '__main__':
    main()